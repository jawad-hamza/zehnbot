import logging
import re
import time
from datetime import datetime, timezone
from typing import List, Optional, Tuple

from fastapi import HTTPException
from sqlalchemy import func, text
from sqlalchemy.orm import Session

from app.config import settings
from app.models.knowledge import KnowledgeChunk
from app.services import embedding_service
from app.services.net_guard import UnsafeURLError, safe_get

logger = logging.getLogger(__name__)

# Big enough to hold a complete answer (a paragraph), small enough that four of them fit the prompt
CHUNK_SIZE = 700
CHUNK_OVERLAP = 100
MAX_URL_BYTES = 2_000_000
CRAWL_TIME_BUDGET_SEC = 120
CRAWL_DELAY_SEC = 0.2
MAX_QUERY_TERMS = 24
CANDIDATES_PER_METHOD = 12   # how many hits keyword and semantic search each contribute before fusion
RRF_K = 60                   # standard reciprocal-rank-fusion damping constant
# Nearest-neighbour search always returns *something*. Beyond this cosine distance the "nearest"
# passage is unrelated to the question, and feeding it to the model only adds noise.
MAX_COSINE_DISTANCE = 0.8

# Only used by the non-Postgres fallback search; Postgres drops stopwords itself.
_STOPWORDS = frozenset(
    "a an and are as at be but by can do does for from has have how i if in is it its me my of on or our "
    "so that the their there they this to was we what when where which who why will with you your".split()
)


class KnowledgeLimitError(Exception):
    pass


def split_into_chunks(text_value: str) -> List[str]:
    """Overlapping ~700 character chunks, cut on a word boundary where one is close."""
    text_value = re.sub(r"\s+", " ", text_value).strip()
    chunks, start = [], 0
    while start < len(text_value):
        end = min(start + CHUNK_SIZE, len(text_value))
        if end < len(text_value):
            space = text_value.rfind(" ", end - 80, end)
            if space > start:
                end = space
        chunks.append(text_value[start:end].strip())
        if end >= len(text_value):
            break
        start = max(end - CHUNK_OVERLAP, start + 1)
    return [c for c in chunks if c]


def replace_source(raw_text: str, client_db_id, source_label: str, db: Session) -> int:
    """Replace the chunks of ONE source label, leaving every other source untouched."""
    db.query(KnowledgeChunk).filter(
        KnowledgeChunk.client_id == client_db_id,
        KnowledgeChunk.source_label == source_label,
    ).delete(synchronize_session=False)
    return _insert_chunks(raw_text, client_db_id, source_label, db)


def delete_knowledge(client_db_id, db: Session, source_label: Optional[str] = None) -> int:
    query = db.query(KnowledgeChunk).filter(KnowledgeChunk.client_id == client_db_id)
    if source_label is not None:
        query = query.filter(KnowledgeChunk.source_label == source_label)
    deleted = query.delete(synchronize_session=False)
    db.commit()
    return deleted


def _insert_chunks(raw_text: str, client_db_id, source_label: str, db: Session) -> int:
    chunks = split_into_chunks(raw_text)
    existing, last_index = (
        db.query(func.count(KnowledgeChunk.id), func.max(KnowledgeChunk.chunk_index))
        .filter(KnowledgeChunk.client_id == client_db_id)
        .one()
    )
    room = settings.MAX_CHUNKS_PER_CLIENT - existing
    if room <= 0:
        db.rollback()
        raise KnowledgeLimitError(f"Knowledge limit reached ({settings.MAX_CHUNKS_PER_CLIENT} chunks). Delete a source first.")
    chunks = chunks[:room]
    start_index = (last_index + 1) if last_index is not None else 0
    rows = [
        KnowledgeChunk(
            client_id=client_db_id,
            chunk_text=chunk,
            chunk_index=start_index + offset,
            source_label=source_label[:255],
        )
        for offset, chunk in enumerate(chunks)
    ]
    db.add_all(rows)
    db.flush()
    chunk_ids = [row.id for row in rows]
    db.commit()   # the text is safe (and keyword-searchable) before we call out to the embeddings API

    if chunks and embedding_service.semantic_search_available(db):
        vectors = embedding_service.embed_documents(chunks)
        if vectors:
            embedding_service.store_embeddings(db, chunk_ids, vectors)
    return len(chunks)


def backfill_embeddings(db: Session, batch_size: int = 200) -> int:
    """Embed chunks that have none yet: knowledge loaded before semantic search was switched on,
    or during an embeddings outage. Used by scripts/embed_backfill.py; safe to re-run."""
    if not embedding_service.semantic_search_available(db):
        return 0
    done = 0
    while True:
        rows = db.execute(
            text("SELECT id, chunk_text FROM knowledge_chunks WHERE embedding IS NULL ORDER BY created_at LIMIT :n"),
            {"n": batch_size},
        ).fetchall()
        if not rows:
            return done
        vectors = embedding_service.embed_documents([r[1] for r in rows])
        if not vectors:
            return done   # provider unavailable; whatever is left is picked up by the next run
        embedding_service.store_embeddings(db, [r[0] for r in rows], vectors)
        done += len(rows)


# ---------- file extraction ----------

def extract_file_text(filename: str, data: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return _extract_pdf(data)
    if lower.endswith(".docx"):
        return _extract_docx(data)
    if lower.endswith((".txt", ".md", ".csv")):
        return data.decode("utf-8", errors="replace")
    raise HTTPException(status_code=400, detail="Unsupported file type. Supported: .pdf, .docx, .txt, .md, .csv")


def _extract_pdf(data: bytes) -> str:
    import io
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
        text_value = "\n\n".join(pages).strip()
    except Exception as e:
        logger.info("PDF parse failed: %s", e)
        raise HTTPException(status_code=400, detail="Could not parse that PDF.")
    if not text_value:
        raise HTTPException(status_code=400, detail="PDF had no extractable text (scanned/image PDFs aren't supported).")
    return text_value


def _extract_docx(data: bytes) -> str:
    import io
    from docx import Document
    try:
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)
        text_value = "\n".join(parts).strip()
    except Exception as e:
        logger.info("DOCX parse failed: %s", e)
        raise HTTPException(status_code=400, detail="Could not parse that DOCX.")
    if not text_value:
        raise HTTPException(status_code=400, detail="DOCX had no extractable text.")
    return text_value


# ---------- web ingestion ----------

def _extract_html_text(content: bytes) -> str:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(content, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "header", "footer", "nav", "aside"]):
        tag.decompose()
    return re.sub(r"\s+", " ", soup.get_text(separator=" ", strip=True)).strip()


def fetch_url_text(url: str) -> str:
    import requests
    try:
        _, status, _, content = safe_get(url.strip(), timeout=15, max_bytes=MAX_URL_BYTES)
    except UnsafeURLError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except requests.RequestException:
        raise HTTPException(status_code=400, detail="Could not fetch that URL.")
    if status != 200:
        raise HTTPException(status_code=400, detail=f"That URL returned HTTP {status}.")
    text_value = _extract_html_text(content)
    if not text_value:
        raise HTTPException(status_code=400, detail="Page had no extractable text")
    return text_value


def crawl_site(start_url: str, max_pages: int = 25) -> List[Tuple[str, str]]:
    """BFS-crawl the same domain as start_url, return list of (url, extracted_text)."""
    import requests
    from bs4 import BeautifulSoup
    from urllib.parse import urljoin, urlparse, urldefrag

    def site_of(u: str) -> str:
        host = (urlparse(u).hostname or "").lower()
        return host[4:] if host.startswith("www.") else host

    start_url = urldefrag(start_url.strip()).url
    domain = site_of(start_url)

    visited = set()
    queue = [start_url]
    results: List[Tuple[str, str]] = []

    deadline = time.monotonic() + CRAWL_TIME_BUDGET_SEC
    skip_exts = (".pdf", ".jpg", ".jpeg", ".png", ".gif", ".svg", ".zip", ".mp4", ".webp", ".ico", ".css", ".js", ".xml", ".woff", ".woff2", ".ttf")

    while queue and len(results) < max_pages and time.monotonic() < deadline:
        url = urldefrag(queue.pop(0)).url
        if url in visited:
            continue
        visited.add(url)

        if urlparse(url).path.lower().endswith(skip_exts):
            continue

        try:
            final_url, status, content_type, content = safe_get(url, timeout=10, max_bytes=MAX_URL_BYTES)
        except (UnsafeURLError, requests.RequestException):
            continue
        if site_of(final_url) != domain:
            if results or len(visited) > 1:
                continue   # a redirect left the site
            domain = site_of(final_url)   # the start URL itself moved (e.g. to a new domain): follow it
        if status != 200 or "html" not in content_type.lower():
            continue

        soup = BeautifulSoup(content, "html.parser")

        # collect same-domain links BEFORE stripping scripts/nav (nav has nav-links too)
        for a in soup.find_all("a", href=True):
            href = urldefrag(urljoin(final_url, a["href"])).url
            p = urlparse(href)
            if p.scheme not in ("http", "https"):
                continue
            if site_of(href) != domain:
                continue
            if href not in visited and href not in queue and len(queue) < 2000:
                queue.append(href)

        text_value = _extract_html_text(content)
        if text_value and len(text_value) > 80:
            results.append((url, text_value))
        time.sleep(CRAWL_DELAY_SEC)

    return results


def run_crawl_job(job_id) -> None:
    """Background task. Owns its own DB session because the request's session is closed by now."""
    from app.database import SessionLocal
    from app.models.job import IngestJob

    db = SessionLocal()
    try:
        job = db.get(IngestJob, job_id)
        if not job:
            return
        job.status = "running"
        client_db_id, url, max_pages = job.client_id, job.url, job.max_pages
        db.commit()

        try:
            pages = crawl_site(url, max_pages=max_pages)
            if not pages:
                raise ValueError("No pages could be crawled from that site.")
            chunks = 0
            note = None
            for page_url, page_text in pages:
                try:
                    chunks += replace_source(page_text, client_db_id, page_url, db)   # re-crawls refresh pages instead of duplicating them
                except KnowledgeLimitError as e:
                    note = str(e)
                    break
            job = db.get(IngestJob, job_id)
            job.status, job.pages_crawled, job.chunks_created, job.error = "done", len(pages), chunks, note
        except Exception as e:
            db.rollback()
            logger.exception("Crawl job %s failed", job_id)
            job = db.get(IngestJob, job_id)
            job.status = "failed"
            job.error = str(e) if isinstance(e, (ValueError, KnowledgeLimitError)) else "The crawl failed unexpectedly."
        job.finished_at = datetime.now(timezone.utc)
        db.commit()
    finally:
        db.close()


# ---------- retrieval ----------

def _terms(query: str) -> List[str]:
    seen, terms = set(), []
    for token in re.findall(r"[^\W_]+", query.lower()):
        if token not in seen:
            seen.add(token)
            terms.append(token)
    return terms[:MAX_QUERY_TERMS]


def search_knowledge(query: str, client_db_id, db: Session, top_k: int = 4) -> List[KnowledgeChunk]:
    """Hybrid retrieval on Postgres: keyword (full-text) and semantic (embedding) search each
    nominate candidates, and reciprocal rank fusion merges the two rankings.

    Keyword search is exact on names, product codes and numbers. Semantic search finds the right
    passage when the visitor uses different words than the page does, or another language.
    Either half works alone, so a missing key or an embeddings outage only lowers quality."""
    terms = _terms(query)
    if db.get_bind().dialect.name != "postgresql":
        return _search_fallback(terms, client_db_id, db, top_k) if terms else []

    rankings = []
    if terms:
        rankings.append(_keyword_ids(terms, client_db_id, db))
    if query.strip() and embedding_service.semantic_search_available(db):
        vector = embedding_service.embed_query(query.strip()[:2000])
        if vector:
            rankings.append(_semantic_ids(vector, client_db_id, db))

    ids = _fuse(rankings)[:top_k]
    if not ids:
        return []
    by_id = {c.id: c for c in db.query(KnowledgeChunk).filter(KnowledgeChunk.id.in_(ids)).all()}
    return [by_id[i] for i in ids if i in by_id]


def _fuse(rankings: List[List]) -> List:
    """Reciprocal rank fusion: score = sum of 1 / (k + rank). Needs no score calibration between
    methods whose raw scores are not comparable (ts_rank vs cosine distance)."""
    scores = {}
    for ranking in rankings:
        for rank, chunk_id in enumerate(ranking):
            scores[chunk_id] = scores.get(chunk_id, 0.0) + 1.0 / (RRF_K + rank + 1)
    return sorted(scores, key=scores.get, reverse=True)


def _keyword_ids(terms: List[str], client_db_id, db: Session) -> List:
    # Indexed full-text search over the generated `search_vector` column (migration 0002).
    # Terms are OR-ed: a visitor's question rarely contains every word of the right passage.
    # `terms` only ever contains letters and digits, so it cannot alter the tsquery syntax.
    rows = db.execute(
        text(
            "SELECT id FROM knowledge_chunks "
            "WHERE client_id = CAST(:cid AS uuid) AND search_vector @@ to_tsquery('english', :q) "
            "ORDER BY ts_rank_cd(search_vector, to_tsquery('english', :q)) DESC "
            "LIMIT :k"
        ),
        {"cid": str(client_db_id), "q": " | ".join(terms), "k": CANDIDATES_PER_METHOD},
    ).fetchall()
    return [r[0] for r in rows]


def _semantic_ids(vector: List[float], client_db_id, db: Session) -> List:
    # Exact nearest-neighbour scan, deliberately without an ANN index: a bot holds at most
    # MAX_CHUNKS_PER_CLIENT vectors, so this takes milliseconds, and an approximate index would
    # apply the tenant filter AFTER picking global neighbours and starve small tenants of results.
    rows = db.execute(
        text(
            "SELECT id FROM knowledge_chunks "
            "WHERE client_id = CAST(:cid AS uuid) AND embedding IS NOT NULL "
            "AND (embedding <=> CAST(:v AS vector)) < :max_distance "
            "ORDER BY embedding <=> CAST(:v AS vector) "
            "LIMIT :k"
        ),
        {
            "cid": str(client_db_id),
            "v": embedding_service.vector_literal(vector),
            "max_distance": MAX_COSINE_DISTANCE,
            "k": CANDIDATES_PER_METHOD,
        },
    ).fetchall()
    return [r[0] for r in rows]


def _search_fallback(terms: List[str], client_db_id, db: Session, top_k: int) -> List[KnowledgeChunk]:
    wanted = {t for t in terms if t not in _STOPWORDS}
    if not wanted:
        return []
    scored = []
    for chunk in db.query(KnowledgeChunk).filter(KnowledgeChunk.client_id == client_db_id).all():
        overlap = len(wanted & set(re.findall(r"[^\W_]+", chunk.chunk_text.lower())))
        if overlap:
            scored.append((overlap, chunk))
    scored.sort(key=lambda x: x[0], reverse=True)
    return [chunk for _, chunk in scored[:top_k]]
